"""File conversions that run on this Mac: documents, data, images, audio, video, and QR codes."""

import csv
import html
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image, ImageOps

DOCS = {".doc", ".docx", ".rtf", ".odt"}
TEXT = {".txt", ".md", ".html", ".htm"}
DATA = {".csv", ".xlsx", ".json"}
SHEETS = {".xls", ".ods"}
SLIDES = {".ppt", ".pptx", ".odp"}
RASTER = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic"}
IMAGES = RASTER | {".svg"}
VIDEO = {".mp4", ".mov", ".m4v", ".mkv", ".avi", ".webm"}
AUDIO = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".ogg"}
MEDIA = VIDEO | AUDIO

IMAGE_OUT = ["png", "jpg", "webp", "gif", "bmp", "tiff", "ico", "pdf"]
AUDIO_OUT = ["mp3", "m4a", "wav"]
PIL_FORMAT = {"png": "PNG", "jpg": "JPEG", "webp": "WEBP", "gif": "GIF", "bmp": "BMP", "tiff": "TIFF", "ico": "ICO", "pdf": "PDF"}
NO_ALPHA = {"jpg", "bmp", "pdf"}
GIF_SECONDS = 20


def same(ext: str, to: str) -> bool:
    return ext.lstrip(".") in {to, {"jpg": "jpeg", "tiff": "tif", "html": "htm"}.get(to)}


def targets(ext: str) -> list[str]:
    ext = ext.lower()
    if ext == ".pdf":
        return ["docx", "txt", "png", "jpg"]
    if ext in DOCS:
        return ["pdf", "csv", "txt"] + ([] if ext == ".docx" else ["docx"])
    if ext in TEXT:
        return ["pdf", "docx"] + (["html"] if ext == ".md" else [])
    if ext in DATA | SHEETS:
        return [t for t in ["csv", "xlsx", "json"] if not same(ext, t)] + (["pdf"] if ext != ".json" else [])
    if ext in SLIDES:
        return ["pdf"]
    if ext in IMAGES:
        return [t for t in IMAGE_OUT if not same(ext, t)]
    if ext in VIDEO:
        return AUDIO_OUT + [t for t in ["mp4", "gif"] if not same(ext, t)]
    if ext in AUDIO:
        return [t for t in AUDIO_OUT if not same(ext, t)]
    return []


FORMATS = {ext: targets(ext) for ext in sorted({".pdf"} | DOCS | TEXT | DATA | SHEETS | SLIDES | IMAGES | MEDIA)}


def convert(name: str, data: bytes, to: str, opt: str = "") -> tuple[str, bytes]:
    """Change format (to="png") or run a tool (to="compress"). Returns (file name, bytes)."""
    src = Path(name)
    ext = src.suffix.lower()
    to = to.lower().lstrip(".")
    stem = src.stem or "converted"
    if to in TOOLS:
        exts, tool, tag = TOOLS[to]
        if ext not in exts:
            raise ValueError(f"This tool doesn't work on {ext or 'these'} files.")
    elif to not in targets(ext):
        raise ValueError(f"Can't convert {ext or 'this file'} to {to}.")
    with tempfile.TemporaryDirectory() as tmp:
        inp = Path(tmp) / f"in{ext}"
        inp.write_bytes(data)
        if to in TOOLS:
            out = tool(inp, opt)
            return f"{stem}-{tag}{out.suffix}", out.read_bytes()
        out = Path(tmp) / f"out.{to}"
        if ext == ".pdf":
            out = {"docx": pdf_to_docx, "txt": pdf_to_txt}.get(to, pdf_to_images)(inp, out, to)
        elif ext in DOCS:
            out = {"pdf": doc_to_pdf, "csv": doc_to_csv, "txt": word_to_text}.get(to, doc_to_docx)(inp, out, to)
        elif ext in TEXT:
            out = text_to(inp, out, to)
        elif ext in SLIDES or (ext in DATA | SHEETS and to == "pdf"):
            out = office_convert(inp, "pdf")
        elif ext in DATA | SHEETS:
            out = write_table(read_table(inp), out)
        elif ext in IMAGES:
            out = convert_image(inp, out, to)
        elif to in AUDIO_OUT:
            out = extract_audio(inp, out, to)
        else:
            out = to_gif(inp, out) if to == "gif" else to_mp4(inp, out)
        return f"{stem}{out.suffix}", out.read_bytes()


# PDF


def pdf_to_docx(inp: Path, out: Path, to: str) -> Path:
    from pdf2docx import Converter

    cv = Converter(str(inp))
    try:
        cv.convert(str(out))
    finally:
        cv.close()
    return out


def pdf_to_txt(inp: Path, out: Path, to: str) -> Path:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(inp))
    try:
        text = "\n\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
    finally:
        pdf.close()
    if not text.strip():
        raise ValueError("No text found. For scanned PDFs, use Image to Text.")
    out.write_text(text.replace("\r\n", "\n"), encoding="utf-8")
    return out


def pdf_to_images(inp: Path, out: Path, to: str) -> Path:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(inp))
    try:
        pages = [pdf[i].render(scale=2).to_pil().convert("RGB") for i in range(len(pdf))]
    finally:
        pdf.close()
    if not pages:
        raise ValueError("PDF has no pages")
    if len(pages) == 1:
        pages[0].save(out, PIL_FORMAT[to])
        return out
    files = []
    for i, page in enumerate(pages, 1):
        buf = io.BytesIO()
        page.save(buf, PIL_FORMAT[to])
        files.append((f"page-{i}.{to}", buf.getvalue()))
    return write_zip(out.with_suffix(".zip"), files)


def write_zip(path: Path, files: list[tuple[str, bytes]]) -> Path:
    with zipfile.ZipFile(path, "w") as z:
        for name, data in files:
            z.writestr(name, data)
    return path


def split_pdf(inp: Path, opt: str) -> Path:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(inp))
    if len(reader.pages) < 2:
        raise ValueError("This PDF has only one page.")
    files = []
    for i, page in enumerate(reader.pages, 1):
        writer = PdfWriter()
        writer.add_page(page)
        buf = io.BytesIO()
        writer.write(buf)
        files.append((f"page-{i}.pdf", buf.getvalue()))
    return write_zip(inp.with_name("pages.zip"), files)


def rotate_pdf(inp: Path, opt: str) -> Path:
    from pypdf import PdfWriter

    angle = int(opt or 90)
    if angle not in (90, 180, 270):
        raise ValueError("Rotate by 90, 180, or 270 degrees.")
    writer = PdfWriter(clone_from=str(inp))
    for page in writer.pages:
        page.rotate(angle)
    out = inp.with_name("rotated.pdf")
    writer.write(out)
    return out


def protect_pdf(inp: Path, opt: str) -> Path:
    from pypdf import PdfWriter

    if not opt:
        raise ValueError("Type a password first.")
    writer = PdfWriter(clone_from=str(inp))
    writer.encrypt(opt, algorithm="AES-256")
    out = inp.with_name("protected.pdf")
    writer.write(out)
    return out


def unlock_pdf(inp: Path, opt: str) -> Path:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(inp))
    if not reader.is_encrypted:
        raise ValueError("This PDF has no password.")
    if not reader.decrypt(opt):
        raise ValueError("Wrong password.")
    writer = PdfWriter(clone_from=reader)
    out = inp.with_name("unlocked.pdf")
    writer.write(out)
    return out


def compress_pdf(inp: Path) -> Path:
    import pymupdf

    doc = pymupdf.open(str(inp))
    try:
        doc.rewrite_images(dpi_threshold=150, dpi_target=120, quality=70)
        out = inp.with_name("small.pdf")
        doc.save(str(out), garbage=4, deflate=True, clean=True)
    finally:
        doc.close()
    return out


def combine(files: list[tuple[str, bytes]]) -> tuple[str, bytes]:
    """Merge PDFs and images, in order, into one PDF."""
    from pypdf import PdfReader, PdfWriter

    if not files:
        raise ValueError("Add some files first.")
    writer = PdfWriter()
    for name, data in files:
        ext = Path(name).suffix.lower()
        if ext == ".pdf":
            writer.append(PdfReader(io.BytesIO(data)))
        elif ext in IMAGES:
            with tempfile.TemporaryDirectory() as tmp:
                inp = Path(tmp) / f"in{ext}"
                inp.write_bytes(data)
                buf = io.BytesIO()
                flatten(open_image(inp)).save(buf, "PDF")
            writer.append(PdfReader(io.BytesIO(buf.getvalue())))
        else:
            raise ValueError(f"{name}: only PDFs and images can be merged.")
    buf = io.BytesIO()
    writer.write(buf)
    return f"{Path(files[0][0]).stem}-merged.pdf", buf.getvalue()


# Word and text


OFFICE_LOCK = threading.Lock()


def soffice() -> str | None:
    for path in (shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if path and Path(path).is_file():
            return path
    return None


def office_profile() -> Path:
    base = Path.home() / ("Library/Caches/U Converter" if sys.platform == "darwin" else ".cache/u-converter")
    return base / "libreoffice"


def office_convert(inp: Path, fmt: str) -> Path:
    """Convert with LibreOffice. fmt can carry a filter, e.g. 'docx:MS Word 2007 XML'."""
    app = soffice()
    if not app:
        raise ValueError("This conversion needs LibreOffice. Install it with: brew install --cask libreoffice")
    outdir = inp.parent / "office"
    outdir.mkdir(exist_ok=True)
    with OFFICE_LOCK:
        run = subprocess.run(
            [app, "--headless", "--norestore", f"-env:UserInstallation={office_profile().as_uri()}",
             "--convert-to", fmt, "--outdir", str(outdir), str(inp)],
            capture_output=True, text=True, timeout=300,
        )
    out = outdir / f"{inp.stem}.{fmt.split(':')[0]}"
    if not out.is_file():
        raise ValueError("LibreOffice could not convert this file. " + (run.stderr.strip().splitlines() or [""])[-1])
    return out


def textutil(inp: Path, fmt: str) -> Path:
    out = inp.with_name(f"textutil.{fmt}")
    subprocess.run(["textutil", "-convert", fmt, str(inp), "-output", str(out)], check=True, capture_output=True)
    return out


def html_to_pdf(page: str, out: Path, base: Path) -> Path:
    from xhtml2pdf import pisa

    with open(out, "wb") as f:
        result = pisa.CreatePDF(page, dest=f, path=str(base))
    if result.err:
        raise ValueError("Could not make a PDF from this document.")
    return out


def doc_to_pdf(inp: Path, out: Path, to: str) -> Path:
    if soffice():
        return office_convert(inp, "pdf")
    page = textutil(inp, "html")
    return html_to_pdf(page.read_text(encoding="utf-8"), out, page)


def doc_to_docx(inp: Path, out: Path, to: str) -> Path:
    return office_convert(inp, "docx") if soffice() else textutil(inp, "docx")


def word_to_text(inp: Path, out: Path, to: str) -> Path:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(inp if inp.suffix == ".docx" else doc_to_docx(inp, out, "docx")))
    lines = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(block, doc)
            listed = para._p.pPr is not None and para._p.pPr.numPr is not None
            lines.append(("• " if listed or para.style.name.startswith("List") else "") + para.text)
        elif tag == "tbl":
            lines.extend("\t".join(c.text.strip() for c in row.cells) for row in Table(block, doc).rows)
            lines.append("")
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out


class Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.tables[-1].append(self.row)
            self.row = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)


def doc_to_csv(inp: Path, out: Path, to: str) -> Path:
    parser = Tables()
    parser.feed(textutil(inp, "html").read_text(encoding="utf-8"))
    tables = [t for t in parser.tables if t]
    if not tables:
        raise ValueError("No tables found in this document. CSV needs a table.")
    with open(out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for i, table in enumerate(tables):
            if i:
                writer.writerow([])
            writer.writerows(table)
    return out


PAGE_STYLE = """<style>
body { font-family: Helvetica; font-size: 11pt; line-height: 1.4; }
table { border-collapse: collapse; } td, th { border: 1px solid #999; padding: 3px 6px; }
pre, code { font-family: Courier; font-size: 9.5pt; background: #f4f4f4; }
</style>"""


def text_html(inp: Path) -> str:
    raw = inp.read_text(encoding="utf-8", errors="replace")
    if inp.suffix in (".html", ".htm"):
        return raw
    if inp.suffix == ".md":
        import markdown

        body = markdown.markdown(raw, extensions=["tables", "fenced_code"])
    else:
        body = "".join(f"<p>{html.escape(line) or '&nbsp;'}</p>" for line in raw.splitlines())
    return f"<html><head><meta charset='utf-8'>{PAGE_STYLE}</head><body>{body}</body></html>"


def text_to(inp: Path, out: Path, to: str) -> Path:
    page = text_html(inp)
    if to == "html":
        out.write_text(page, encoding="utf-8")
        return out
    source = inp if inp.suffix in (".html", ".htm") else inp.with_name("page.html")
    if source is not inp:
        source.write_text(page, encoding="utf-8")
    if soffice():
        return office_convert(source, "pdf:writer_web_pdf_Export" if to == "pdf" else "docx:MS Word 2007 XML")
    return html_to_pdf(page, out, inp) if to == "pdf" else textutil(source, "docx")


# Spreadsheets and data


def cell(value):
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def read_table(inp: Path) -> list[list]:
    if inp.suffix == ".csv":
        with open(inp, newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))
    if inp.suffix in SHEETS:
        inp = office_convert(inp, "xlsx")
    if inp.suffix == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(inp, read_only=True, data_only=True).active
        return [["" if v is None else v for v in row] for row in sheet.iter_rows(values_only=True)]
    data = json.loads(inp.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [data])
    if not isinstance(data, list):
        raise ValueError("JSON must be a list of rows.")
    if data and all(isinstance(r, dict) for r in data):
        header = list(dict.fromkeys(k for r in data for k in r))
        return [header] + [[cell(r.get(k, "")) for k in header] for r in data]
    return [[cell(v) for v in r] if isinstance(r, list) else [cell(r)] for r in data]


def number(value):
    if isinstance(value, str) and re.fullmatch(r"-?\d+(\.\d+)?", value.strip()):
        return float(value) if "." in value else int(value)
    return value


def write_table(rows: list[list], out: Path) -> Path:
    if not rows:
        raise ValueError("This file has no rows.")
    if out.suffix == ".csv":
        with open(out, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)
    elif out.suffix == ".xlsx":
        from openpyxl import Workbook

        book = Workbook()
        for row in rows:
            book.active.append([number(v) for v in row])
        book.save(out)
    else:
        header = [str(h) for h in rows[0]]
        records = [dict(zip(header, map(number, row))) for row in rows[1:]]
        out.write_text(json.dumps(records, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    return out


# Images


def open_image(inp: Path) -> Image.Image:
    if inp.suffix == ".svg":
        import pymupdf

        doc = pymupdf.open(str(inp))
        pix = doc[0].get_pixmap(dpi=192, alpha=True)
        doc.close()
        return Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
    if inp.suffix == ".heic":
        from pillow_heif import register_heif_opener

        register_heif_opener()
    return ImageOps.exif_transpose(Image.open(inp))


def flatten(im: Image.Image) -> Image.Image:
    if im.mode == "RGB":
        return im
    rgba = im.convert("RGBA")
    flat = Image.new("RGB", rgba.size, "white")
    flat.paste(rgba, mask=rgba.getchannel("A"))
    return flat


def save_image(im: Image.Image, out: Path, to: str, **options) -> Path:
    if to in NO_ALPHA:
        im = flatten(im)
    if to == "ico":
        side = max(im.size)
        square = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        square.paste(im.convert("RGBA"), ((side - im.width) // 2, (side - im.height) // 2))
        im = square
        options["sizes"] = [(s, s) for s in (16, 32, 48, 64, 128, 256) if s <= side] or [(side, side)]
    im.save(out, PIL_FORMAT[to], **options)
    return out


def convert_image(inp: Path, out: Path, to: str) -> Path:
    return save_image(open_image(inp), out, to)


def image_kind(inp: Path) -> str:
    kind = {"jpeg": "jpg", "tif": "tiff"}.get(inp.suffix[1:], inp.suffix[1:])
    return kind if kind in PIL_FORMAT else "jpg"


def compress_image(inp: Path) -> Path:
    im = open_image(inp)
    kind = image_kind(inp)
    kind = kind if kind in ("png", "jpg", "webp") else "jpg"
    out = inp.with_name(f"small.{kind}")
    if kind == "png":
        im.convert("RGBA").quantize(256, method=Image.Quantize.FASTOCTREE).save(out, "PNG", optimize=True)
        return out
    return save_image(im, out, kind, quality=70, optimize=True)


def resize_image(inp: Path, opt: str) -> Path:
    im = open_image(inp)
    opt = opt or "50%"
    if opt.endswith("%"):
        scale = int(opt[:-1]) / 100
    else:
        scale = min(1.0, int(opt.rstrip("px")) / max(im.size))
    size = (max(1, round(im.width * scale)), max(1, round(im.height * scale)))
    kind = image_kind(inp)
    out = inp.with_name(f"resized.{kind}")
    return save_image(im.resize(size, Image.Resampling.LANCZOS), out, kind, **({"quality": 90} if kind == "jpg" else {}))


def strip_metadata(inp: Path, opt: str) -> Path:
    im = open_image(inp)
    clean = Image.frombytes(im.mode, im.size, im.tobytes())
    if im.mode == "P":
        clean.putpalette(im.getpalette())
    kind = image_kind(inp)
    out = inp.with_name(f"clean.{kind}")
    return save_image(clean, out, kind, **({"quality": 95} if kind == "jpg" else {}))


def read_qr(name: str, data: bytes) -> str:
    import cv2
    import numpy as np

    ext = Path(name).suffix.lower()
    if ext not in IMAGES:
        raise ValueError("Read QR Code needs an image.")
    with tempfile.TemporaryDirectory() as tmp:
        inp = Path(tmp) / f"in{ext}"
        inp.write_bytes(data)
        pixels = cv2.cvtColor(np.array(flatten(open_image(inp))), cv2.COLOR_RGB2BGR)
    detector = cv2.QRCodeDetector()
    ok, texts, _, _ = detector.detectAndDecodeMulti(pixels)
    found = [t for t in texts if t] if ok else []
    if not found:
        text, _, _ = detector.detectAndDecode(pixels)
        found = [text] if text else []
    if not found:
        raise ValueError("No QR code found in this image.")
    return "\n".join(found)


# Audio and video


def ffmpeg() -> str:
    for path in (shutil.which("ffmpeg"), "/opt/homebrew/bin/ffmpeg", "/usr/local/bin/ffmpeg"):
        if path and Path(path).is_file():
            return path
    raise ValueError("Audio and video need ffmpeg. Install it with: brew install ffmpeg")


def run_ffmpeg(inp: Path, out: Path, *args: str) -> Path:
    run = subprocess.run([ffmpeg(), "-y", "-i", str(inp), *args, str(out)], capture_output=True, text=True)
    if run.returncode != 0:
        last = run.stderr.strip().splitlines()[-1:] or ["unknown error"]
        raise ValueError(f"ffmpeg failed: {last[0]}")
    return out


def extract_audio(inp: Path, out: Path, to: str) -> Path:
    codec = {"mp3": ["-c:a", "libmp3lame", "-q:a", "2"], "m4a": ["-c:a", "aac", "-b:a", "192k"], "wav": ["-c:a", "pcm_s16le"]}[to]
    return run_ffmpeg(inp, out, "-vn", *codec)


def to_mp4(inp: Path, out: Path, crf: str = "20", longest: int = 0) -> Path:
    scale = ["-vf", f"scale='if(gt(iw,ih),min({longest},iw),-2)':'if(gt(iw,ih),-2,min({longest},ih))'"] if longest else []
    return run_ffmpeg(
        inp, out, *scale, "-c:v", "libx264", "-crf", crf, "-preset", "veryfast", "-pix_fmt", "yuv420p",
        "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart",
    )


def to_gif(inp: Path, out: Path) -> Path:
    graph = "fps=12,scale='min(480,iw)':-2:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse"
    return run_ffmpeg(inp, out, "-t", str(GIF_SECONDS), "-vf", graph, "-loop", "0")


def compress_video(inp: Path) -> Path:
    return to_mp4(inp, inp.with_name("small.mp4"), crf="28", longest=1280)


def mute_video(inp: Path, opt: str) -> Path:
    return run_ffmpeg(inp, inp.with_name(f"muted{inp.suffix}"), "-c:v", "copy", "-an")


def compress(inp: Path, opt: str) -> Path:
    if inp.suffix == ".pdf":
        out = compress_pdf(inp)
    elif inp.suffix in VIDEO:
        out = compress_video(inp)
    else:
        out = compress_image(inp)
    if out.stat().st_size >= inp.stat().st_size and out.suffix == inp.suffix:
        raise ValueError("This file is already as small as it gets.")
    return out


TOOLS = {
    "compress": ({".pdf"} | RASTER | VIDEO, compress, "compressed"),
    "split": ({".pdf"}, split_pdf, "pages"),
    "rotate": ({".pdf"}, rotate_pdf, "rotated"),
    "protect": ({".pdf"}, protect_pdf, "protected"),
    "unlock": ({".pdf"}, unlock_pdf, "unlocked"),
    "resize": (RASTER, resize_image, "resized"),
    "strip": (RASTER, strip_metadata, "clean"),
    "mute": (VIDEO, mute_video, "muted"),
}


def qr_code(text: str, kind: str = "png") -> bytes:
    import segno

    if not text.strip():
        raise ValueError("Type some text or a link first.")
    buf = io.BytesIO()
    segno.make_qr(text, error="m").save(buf, kind=kind, scale=10, border=4)
    return buf.getvalue()
